"""One loader per MultiMedQA dataset -> list[Record] (unified schema).

Every transformation is labelled. Nothing is invented: where the paper is silent
the choice is tagged REPLICATION ASSUMPTION and cross-referenced in docs/deviations.md.
"""
from __future__ import annotations

import csv
import io as _io
import json
from pathlib import Path

import openpyxl
from lxml import etree

from ..utils import paths
from .schema import Record

LETTERS = ["A", "B", "C", "D", "E"]


def _jsonl(p: Path) -> list[dict]:
    # NB: iterate the file (or split on "\n") -- NOT str.splitlines(), which also
    # breaks on U+2028/U+2029/U+000B etc. that json.dumps(ensure_ascii=False) may
    # leave literal inside string values.
    with open(p, "r", encoding="utf-8") as fh:
        return [json.loads(line) for line in fh if line.strip()]


# --------------------------------------------------------------------------- #
# MedQA (USMLE), 4-option                        PAPER-SPECIFIED (Section 3.1)
# --------------------------------------------------------------------------- #
def load_medqa() -> list[Record]:
    base = paths.RAW_MEDQA / "US" / "4_options"
    out: list[Record] = []
    for split in ("train", "dev", "test"):
        for i, r in enumerate(_jsonl(base / f"phrases_no_exclude_{split}.jsonl")):
            out.append(Record(
                uid=f"medqa_usmle_4opt/{split}/{i}",
                dataset="medqa_usmle_4opt", split=split, task="mc",
                question=r["question"],
                options={k: r["options"][k] for k in ("A", "B", "C", "D")},
                gold=r["answer_idx"],                       # DATASET-VERIFIED
                gold_source="answer_idx (dataset field)",
                meta={"usmle_step": r.get("meta_info")},
            ))
    return out


# --------------------------------------------------------------------------- #
# MedMCQA                                        PAPER-SPECIFIED (Section 3.1)
# --------------------------------------------------------------------------- #
def load_medmcqa() -> list[Record]:
    out: list[Record] = []
    for split in ("train", "dev", "test"):     # 'dev' file == official validation split
        norm_split = "validation" if split == "dev" else split
        for i, r in enumerate(_jsonl(paths.RAW_MEDMCQA / f"{split}.json")):
            cop = r.get("cop")
            # DATASET-VERIFIED: cop is 1-indexed (1=opa..4=opd); test split has none.
            # NB: MedMCQA `id` values are NOT globally unique -> index the uid, keep id in meta.
            gold = LETTERS[cop - 1] if isinstance(cop, int) and cop in (1, 2, 3, 4) else None
            out.append(Record(
                uid=f"medmcqa/{norm_split}/{i}",
                dataset="medmcqa", split=norm_split, task="mc",
                question=r["question"],
                options={"A": r["opa"], "B": r["opb"], "C": r["opc"], "D": r["opd"]},
                gold=gold,
                gold_source="cop (1-indexed) -> letter" if gold else "withheld (test split)",
                meta={"medmcqa_id": r.get("id"),
                      "subject": r.get("subject_name"), "topic": r.get("topic_name"),
                      "choice_type": r.get("choice_type"), "explanation": r.get("exp")},
            ))
    return out


# --------------------------------------------------------------------------- #
# PubMedQA (pqa_labeled)                          PAPER-SPECIFIED (Section 3.1)
# closed-domain: question + abstract context -> yes/no/maybe
# --------------------------------------------------------------------------- #
PUBMEDQA_OPTIONS = {"A": "yes", "B": "no", "C": "maybe"}   # PAPER-SPECIFIED (Table A.15 prompt)
_PUBMEDQA_LETTER = {"yes": "A", "no": "B", "maybe": "C"}


def load_pubmedqa() -> list[Record]:
    data = json.loads((paths.RAW_PUBMEDQA / "ori_pqal.json").read_text(encoding="utf-8"))
    test_gt = json.loads((paths.RAW_PUBMEDQA / "test_ground_truth.json").read_text(encoding="utf-8"))
    test_ids = set(test_gt)                                # RA-02: official 500-item test split
    out: list[Record] = []
    for pmid, r in data.items():
        split = "test" if pmid in test_ids else "train"    # RA-02: remaining 500 -> exemplar pool
        ctx = r["CONTEXTS"]
        context = "\n".join(ctx) if isinstance(ctx, list) else str(ctx)
        decision = r["final_decision"]
        out.append(Record(
            uid=f"pubmedqa/{split}/{pmid}",
            dataset="pubmedqa", split=split, task="mc",
            question=r["QUESTION"], options=dict(PUBMEDQA_OPTIONS), context=context,
            gold=_PUBMEDQA_LETTER.get(decision),
            gold_source="final_decision -> letter",
            meta={"pmid": pmid, "year": r.get("YEAR"), "long_answer": r.get("LONG_ANSWER")},
        ))
    return out


# --------------------------------------------------------------------------- #
# MMLU — 6 clinical subjects                      PAPER-SPECIFIED (Section 3.1)
# headerless CSV: question, A, B, C, D, answer_letter
# --------------------------------------------------------------------------- #
def load_mmlu() -> list[Record]:
    out: list[Record] = []
    for subject in paths.MMLU_SUBJECTS:
        for split in ("dev", "validation", "test"):
            p = paths.RAW_MMLU / subject / f"{split}.csv"
            if not p.exists():
                continue
            rows = [r for r in csv.reader(_io.StringIO(p.read_text(encoding="utf-8"))) if r]
            for i, r in enumerate(rows):
                q, a, b, c, d, ans = r
                out.append(Record(
                    uid=f"mmlu_{subject}/{split}/{i}",
                    dataset=f"mmlu_{subject}", split=split, task="mc",
                    question=q, options={"A": a, "B": b, "C": c, "D": d},
                    gold=ans.strip(), gold_source="csv column 5",
                    meta={"mmlu_subject": subject},
                ))
    return out


# --------------------------------------------------------------------------- #
# LiveQA TREC-2017 (Medical) — long-form          PAPER-SPECIFIED (Section 3.1)
# --------------------------------------------------------------------------- #
def _txt(el) -> str:
    return " ".join("".join(el.itertext()).split()) if el is not None else ""


def load_liveqa() -> list[Record]:
    out: list[Record] = []
    # --- test split: canonical NIST XML ---
    test_xml = paths.RAW_LIVEQA / "TestDataset" / "TREC-2017-LiveQA-Medical-Test.xml"
    root = etree.fromstring(test_xml.read_bytes())
    for q in root.findall(".//NLM-QUESTION"):
        qid = q.get("qid")
        subj = _txt(q.find(".//SUBJECT"))
        msg = _txt(q.find(".//MESSAGE"))
        question = f"{subj}: {msg}" if subj else msg
        # NB: the test XML mixes <RefAnswer> and <ReferenceAnswer> child tags.
        ans_nodes = q.findall(".//RefAnswer") + q.findall(".//ReferenceAnswer")
        answers = [a for a in (_txt(n.find("ANSWER")) for n in ans_nodes) if a]
        out.append(Record(
            uid=f"liveqa/test/{qid}", dataset="liveqa", split="test", task="longform",
            question=question,
            gold=answers[0] if answers else None,
            gold_source="NIST ReferenceAnswers (not used for automated scoring)",
            meta={"qid": qid, "subject": subj, "message": msg,
                  "nist_paraphrase": _txt(q.find(".//NIST-PARAPHRASE")),
                  "n_reference_answers": len(answers), "all_reference_answers": answers},
        ))
    # --- train split: community mirror (RA-20, UNVERIFIED provenance) ---
    tp = paths.RAW_LIVEQA / "TrainingDatasets" / "train_pairs.jsonl"
    if tp.exists():
        for i, r in enumerate(_jsonl(tp)):
            subj = (r.get("subject") or "").strip()
            msg = (r.get("message") or "").strip()
            question = f"{subj}: {msg}" if subj else msg
            out.append(Record(
                uid=f"liveqa/train/{i}",
                dataset="liveqa", split="train", task="longform",
                question=question, gold=(r.get("answer") or "").strip() or None,
                gold_source="truehealth/liveqa mirror (RA-20, UNVERIFIED)",
                meta={"questionid": r.get("questionid"), "focus": r.get("focus"),
                      "type": r.get("type"), "pairid": r.get("pairid")},
            ))
    return out


# --------------------------------------------------------------------------- #
# MedicationQA (MedInfo 2019) — long-form         PAPER-SPECIFIED (Section 3.1)
# --------------------------------------------------------------------------- #
def load_medicationqa() -> list[Record]:
    p = paths.RAW_MEDICATIONQA / "MedInfo2019-QA-Medications.xlsx"
    ws = openpyxl.load_workbook(p, read_only=True)["DrugQA"]
    rows = list(ws.iter_rows(values_only=True))[1:]
    out: list[Record] = []
    for i, r in enumerate(rows):
        q, focus, qtype, ans, section, url = (list(r) + [None] * 6)[:6]
        if not q or not str(q).strip():
            continue
        out.append(Record(
            uid=f"medicationqa/all/{i}", dataset="medicationqa", split="all", task="longform",
            question=str(q).strip(),
            gold=str(ans).strip() if ans else None,
            gold_source="Answer column (not used for automated scoring)",
            meta={"focus_drug": focus, "question_type": qtype,
                  "section_title": section, "url": url},
        ))
    return out


# --------------------------------------------------------------------------- #
# HealthSearchQA (this paper's dataset) — long-form, question-only
# --------------------------------------------------------------------------- #
def load_healthsearchqa() -> list[Record]:
    p = paths.RAW_HEALTHSEARCHQA / "41586_2023_6291_MOESM6_ESM.xlsx"
    wb = openpyxl.load_workbook(p, read_only=True)
    out: list[Record] = []
    for i, row in enumerate(wb.worksheets[0].iter_rows(values_only=True)):
        if not row or not row[0] or not str(row[0]).strip():
            continue
        out.append(Record(
            uid=f"healthsearchqa/all/{i}", dataset="healthsearchqa", split="all",
            task="longform", question=str(row[0]).strip(),
            gold=None, gold_source="NOT APPLICABLE (question-only dataset)",
            meta={},
        ))
    return out


LOADERS = {
    "medqa_usmle_4opt": load_medqa,
    "medmcqa": load_medmcqa,
    "pubmedqa": load_pubmedqa,
    "mmlu": load_mmlu,
    "liveqa": load_liveqa,
    "medicationqa": load_medicationqa,
    "healthsearchqa": load_healthsearchqa,
}
